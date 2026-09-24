"""
文书导出服务（DOCX 红头笔录生成）

实现需求文档：
- FR-3.5.1 红头笔录在线实时预览；
- FR-3.5.2 DOCX 标准文书排版生成与下载；
- 8.1 标准询问笔录文书排版规格（红头标题、案件信息、问答正文、落款）；
- 8.2 五流证据要素覆盖表、8.3 AI 侦查研判报告（可选附带）；
- IR-2 空白过滤、IR-3 红头排版、IR-4 版本指纹。

排版规范（8.1）：
- 红头标题：红色字体、居中；
- 问答正文："问："加粗（黑体），"答："正体（宋体），标准行距；
- 未作答空白项自动过滤（IR-2）。
"""

import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.enums import (
    CHAPTER_LABEL,
    CHAPTER_ORDER,
    AuditOpType,
    Chapter,
    ElementStatus,
    FlowType,
)
from app.core.exceptions import NotFoundError, ValidationError
from app.core.security import compute_fingerprint
from app.models.record import QAItem, RecordSession
from app.models.user import User
from app.schemas.document import (
    DocumentExportRequest,
    DocumentPreview,
    DocumentPreviewItem,
)
from app.schemas.fiveflow import FiveFlowAnalysisResult
from app.services.audit_service import AuditService
from app.services.fiveflow_service import FiveFlowService

# 中文字体常量（8.1 字体规范：正文宋体，标题黑体）
_FONT_SONG = "宋体"
_FONT_HEI = "黑体"
# 红头标题红色
_RED_HEADER_COLOR = RGBColor(0xC0, 0x00, 0x00)


class DocumentService:
    """文书预览与 DOCX 生成服务。"""

    def __init__(self, db: Session):
        self.db = db
        self.audit = AuditService(db)
        self.fiveflow_service = FiveFlowService(db)

    # ============================================================
    # 一、预览数据构建（FR-3.5.1）
    # ============================================================
    def build_preview(self, session_id: str) -> DocumentPreview:
        """构建红头笔录预览数据（FR-3.5.1 处理逻辑）。

        读取会话已作答问答项 → 按公安询问笔录排版规范构建预览文档流 → 过滤未作答空白项。
        """
        session = self._get_session(session_id)
        case = session.case
        if case is None:
            raise NotFoundError("案件信息不存在")

        # 读取已作答问答项（IR-2 空白过滤），按章节+排序组织
        qa_items = self.db.execute(
            select(QAItem)
            .where(QAItem.session_id == session_id)
            .order_by(QAItem.sort_order)
        ).scalars().all()
        answered = [q for q in qa_items if q.is_answered]

        paragraphs: list[DocumentPreviewItem] = []
        current_chapter = None
        for q in answered:
            chapter_label = None
            if q.chapter != current_chapter:
                current_chapter = q.chapter
                try:
                    chapter_label = CHAPTER_LABEL[Chapter(q.chapter)]
                except (ValueError, KeyError):
                    chapter_label = None
            paragraphs.append(DocumentPreviewItem(
                chapter_label=chapter_label,
                question=q.question,
                answer=q.answer or "",
            ))

        # 询问人（办案民警姓名）
        creator = self.db.execute(select(User).where(User.officer_no == session.creator_no)).scalar_one_or_none()
        inquirer = creator.name if creator else session.creator_no
        interviewee = case.reporter.name if case.reporter else None

        red_header_title = f"{case.handling_org or '公安机关'}询问笔录"

        return DocumentPreview(
            session_id=session_id,
            red_header_title=red_header_title,
            case_no=case.case_no,
            case_category=case.case_category,
            inquiry_time=case.report_time.strftime("%Y年%m月%d日 %H:%M") if case.report_time else None,
            inquiry_place=case.handling_org,
            inquirer=inquirer,
            interviewee=interviewee,
            qa_paragraphs=paragraphs,
            answered_count=len(answered),
            signature_area={
                "interviewee_sign": "被询问人（核对无误后签名）：____________",
                "officer_sign": "办案人员（署名）：____________",
                "date": "日期：______年____月____日",
            },
        )

    # ============================================================
    # 二、DOCX 生成（FR-3.5.2 / 8.1 排版规格）
    # ============================================================
    def generate_docx(
        self,
        session_id: str,
        options: DocumentExportRequest,
        user: User,
        request=None,
    ) -> tuple[bytes, str, str]:
        """生成标准排版 DOCX 文书（FR-3.5.2 处理逻辑）。

        步骤：
        1) 读取已作答问答与案件主数据；
        2) 按红头排版标准（宋体、加粗"问："标记、标准行距）构建文档流；
        3) 过滤未作答空白项（IR-2）；
        4) 可选附带五流覆盖表（8.2）与 AI 研判报告（8.3）；
        5) 生成文书版本指纹并写入审计（IR-4）。

        :return: (DOCX 二进制内容, 下载文件名, 版本指纹)
        """
        preview = self.build_preview(session_id)
        if preview.answered_count == 0:
            # 无已作答内容提示（FR-3.5.1 异常与边界）
            raise ValidationError("无已作答的问答内容，无法生成文书")

        doc = Document()
        self._set_default_font(doc)

        # --- 红头标题（8.1：红色、居中）---
        self._add_red_header(doc, preview.red_header_title)

        # --- 案件基本信息栏（8.1：宋体，标准信息栏排版）---
        self._add_case_info_block(doc, preview)

        # --- 问答正文（8.1：分章节问答对，"问："加粗黑体，"答："正体宋体）---
        self._add_qa_body(doc, preview.qa_paragraphs)

        # --- 落款信息（8.1：签名区，预留签字位置）---
        self._add_signature_block(doc, preview.signature_area)

        # --- 可选：五流证据覆盖表（8.2）---
        if options.include_five_flow_table:
            session = self._get_session(session_id)
            analysis = self.fiveflow_service.compute_coverage(session)
            self._add_five_flow_table(doc, analysis)

        # --- 可选：AI 侦查研判报告（8.3）---
        if options.include_analysis_report:
            self._add_analysis_report(doc, session_id)

        # 序列化为二进制
        buffer = io.BytesIO()
        doc.save(buffer)
        content = buffer.getvalue()

        # 版本指纹（IR-4：导出文书生成版本指纹并写入审计）
        fingerprint = compute_fingerprint(content)
        file_name = f"询问笔录_{preview.case_no}_{datetime.now().strftime('%Y%m%d%H%M%S')}.docx"

        # 文书导出审计（2.4.2：导出时间、文书版本指纹）
        self.audit.log(
            op_type=AuditOpType.DOCUMENT_EXPORT,
            officer_no=user.officer_no,
            case_no=preview.case_no,
            session_id=session_id,
            request=request,
            description=f"导出 DOCX 笔录：{file_name}",
            detail={
                "file_name": file_name, "answered_count": preview.answered_count,
                "include_five_flow": options.include_five_flow_table,
                "include_report": options.include_analysis_report,
            },
            change_fingerprint=fingerprint,
        )
        return content, file_name, fingerprint

    # ---------- DOCX 排版辅助方法（8.1 排版规格） ----------
    def _set_default_font(self, doc: Document) -> None:
        """设置文档默认字体为宋体（8.1 字体规范）。"""
        style = doc.styles["Normal"]
        style.font.name = _FONT_SONG
        style.font.size = Pt(12)
        # 设置中文字体（需通过 rPr 的 eastAsia 属性）
        style.element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_SONG)

    def _add_red_header(self, doc: Document, title: str) -> None:
        """添加红头标题（8.1：红色文印标题，居中，红色字体）。"""
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.font.name = _FONT_HEI
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = _RED_HEADER_COLOR
        run._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_HEI)
        # 红头下方分隔线（空段落）
        doc.add_paragraph()

    def _add_case_info_block(self, doc: Document, preview: DocumentPreview) -> None:
        """添加案件基本信息栏（8.1：时间、地点、询问人、被询问人，宋体标准信息栏）。"""
        info_lines = [
            f"案件编号：{preview.case_no}",
            f"案件类别：{preview.case_category or ''}",
            f"询问时间：{preview.inquiry_time or ''}",
            f"询问地点：{preview.inquiry_place or ''}",
            f"询 问 人：{preview.inquirer or ''}",
            f"被询问人：{preview.interviewee or ''}",
        ]
        for line in info_lines:
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.name = _FONT_SONG
            run.font.size = Pt(12)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_SONG)
        doc.add_paragraph()  # 信息栏与正文间空行

    def _add_qa_body(self, doc: Document, paragraphs: list[DocumentPreviewItem]) -> None:
        """添加问答正文（8.1：分章节；"问："加粗黑体，"答："正体宋体；标准行距）。"""
        for item in paragraphs:
            # 章节标题（若切换章节）
            if item.chapter_label:
                cp = doc.add_paragraph()
                crun = cp.add_run(f"【{item.chapter_label}】")
                crun.font.name = _FONT_HEI
                crun.font.size = Pt(12)
                crun.font.bold = True
                crun._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_HEI)

            # 问：加粗黑体
            qp = doc.add_paragraph()
            qp.paragraph_format.line_spacing = 1.5  # 标准行距
            q_prefix = qp.add_run("问：")
            q_prefix.font.name = _FONT_HEI
            q_prefix.font.bold = True
            q_prefix.font.size = Pt(12)
            q_prefix._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_HEI)
            q_body = qp.add_run(item.question)
            q_body.font.name = _FONT_HEI
            q_body.font.bold = True
            q_body.font.size = Pt(12)
            q_body._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_HEI)

            # 答：正体宋体
            ap = doc.add_paragraph()
            ap.paragraph_format.line_spacing = 1.5
            a_prefix = ap.add_run("答：")
            a_prefix.font.name = _FONT_SONG
            a_prefix.font.size = Pt(12)
            a_prefix._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_SONG)
            a_body = ap.add_run(item.answer)
            a_body.font.name = _FONT_SONG
            a_body.font.size = Pt(12)
            a_body._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_SONG)

    def _add_signature_block(self, doc: Document, signature: dict) -> None:
        """添加落款信息（8.1：被询问人核对签名区、办案人员署名区，预留签字位置）。"""
        doc.add_paragraph()
        for key in ("interviewee_sign", "officer_sign", "date"):
            if key in signature:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                run = p.add_run(signature[key])
                run.font.name = _FONT_SONG
                run.font.size = Pt(12)
                run._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_SONG)

    def _add_five_flow_table(self, doc: Document, analysis: FiveFlowAnalysisResult) -> None:
        """添加五流证据要素覆盖表（8.2 格式）。"""
        doc.add_page_break()
        title = doc.add_paragraph()
        trun = title.add_run("五流证据要素覆盖表")
        trun.font.name = _FONT_HEI
        trun.font.bold = True
        trun.font.size = Pt(14)
        trun._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_HEI)

        # 表头：流类型 | 收集状态 | 已提取证据摘要 | 重点缺口清单
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        headers = ["流类型", "收集状态", "已提取证据摘要", "重点缺口清单"]
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = h

        status_label = {
            ElementStatus.COLLECTED: "已收集",
            ElementStatus.NEED_SUPPLEMENT: "需补充",
            ElementStatus.KEY_GAP: "重点缺口",
            ElementStatus.NOT_INVOLVED: "不涉及",
        }
        for flow in analysis.flows:
            row = table.add_row().cells
            row[0].text = flow.flow_label
            row[1].text = f"{status_label[flow.status]}（{flow.coverage}%）"
            row[2].text = "；".join(flow.evidence_summary) if flow.evidence_summary else "—"
            row[3].text = "、".join(flow.key_gaps) if flow.key_gaps else "—"

    def _add_analysis_report(self, doc: Document, session_id: str) -> None:
        """添加 AI 侦查研判指引报告（8.3 格式）。"""
        from app.models.record import AiSuggestion

        doc.add_page_break()
        title = doc.add_paragraph()
        trun = title.add_run("AI 侦查研判指引报告")
        trun.font.name = _FONT_HEI
        trun.font.bold = True
        trun.font.size = Pt(14)
        trun._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_HEI)

        suggestions = self.db.execute(
            select(AiSuggestion).where(AiSuggestion.session_id == session_id)
        ).scalars().all()
        if not suggestions:
            doc.add_paragraph("（暂无研判建议）")
            return
        for s in suggestions:
            p = doc.add_paragraph()
            run = p.add_run(f"{s.title or ''} {s.content}")
            run.font.name = _FONT_SONG
            run.font.size = Pt(12)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_SONG)

    def _get_session(self, session_id: str) -> RecordSession:
        """获取会话（含案件）。"""
        from sqlalchemy.orm import selectinload
        from app.models.case import Case

        session = self.db.execute(
            select(RecordSession)
            .options(selectinload(RecordSession.case).selectinload(Case.reporter))
            .where(RecordSession.id == session_id)
        ).scalar_one_or_none()
        if session is None:
            raise NotFoundError("笔录会话不存在")
        return session
